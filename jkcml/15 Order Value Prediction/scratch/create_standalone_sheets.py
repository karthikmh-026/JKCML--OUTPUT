import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load the master workbook
src_wb_path = "ML_15_Order_Value_Prediction_Enriched_Output_Clean.xlsx"
print(f"Loading {src_wb_path}...")
src_wb = openpyxl.load_workbook(src_wb_path, data_only=True)

# 1. Standalone Feature_Importance.xlsx
fi_wb = openpyxl.Workbook()
fi_ws = fi_wb.active
fi_ws.title = "Feature_Importance"
src_fi = src_wb["Feature_Importance"]

for row in src_fi.iter_rows(values_only=False):
    for cell in row:
        new_cell = fi_ws.cell(row=cell.row, column=cell.column, value=cell.value)
        if cell.has_style:
            new_cell.font = Font(name=cell.font.name, size=cell.font.size, bold=cell.font.bold, color=cell.font.color)
            new_cell.fill = PatternFill(fill_type=cell.fill.fill_type, start_color=cell.fill.start_color, end_color=cell.fill.end_color)
            new_cell.alignment = Alignment(horizontal=cell.alignment.horizontal, vertical=cell.alignment.vertical)
            new_cell.border = Border(
                left=cell.border.left, right=cell.border.right,
                top=cell.border.top, bottom=cell.border.bottom
            )
            if cell.number_format:
                new_cell.number_format = cell.number_format

fi_ws.column_dimensions['A'].width = 34
fi_ws.column_dimensions['B'].width = 18
fi_ws.column_dimensions['C'].width = 90
fi_wb.save("ML_15_Feature_Importance.xlsx")
print("Saved ML_15_Feature_Importance.xlsx")

# 2. Standalone Model_Ranking_Evaluation.xlsx
mr_wb = openpyxl.Workbook()
mr_ws = mr_wb.active
mr_ws.title = "Model Ranking & Evaluation"
src_mr = src_wb["Model Ranking & Evaluation"]

for row in src_mr.iter_rows(values_only=False):
    for cell in row:
        new_cell = mr_ws.cell(row=cell.row, column=cell.column, value=cell.value)
        if cell.has_style:
            new_cell.font = Font(name=cell.font.name, size=cell.font.size, bold=cell.font.bold, color=cell.font.color)
            new_cell.fill = PatternFill(fill_type=cell.fill.fill_type, start_color=cell.fill.start_color, end_color=cell.fill.end_color)
            new_cell.alignment = Alignment(horizontal=cell.alignment.horizontal, vertical=cell.alignment.vertical)
            new_cell.border = Border(
                left=cell.border.left, right=cell.border.right,
                top=cell.border.top, bottom=cell.border.bottom
            )
            if cell.number_format:
                new_cell.number_format = cell.number_format

for col in mr_ws.columns:
    max_len = max(len(str(c.value or '')) for c in col)
    col_letter = get_column_letter(col[0].column)
    mr_ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

mr_wb.save("ML_15_Model_Ranking_Evaluation.xlsx")
print("Saved ML_15_Model_Ranking_Evaluation.xlsx")

# 3. Standalone Summary_KPI_Dashboard.xlsx
kpi_wb = openpyxl.Workbook()
kpi_ws = kpi_wb.active
kpi_ws.title = "Summary KPI Dashboard"
src_kpi = src_wb["Summary KPI Dashboard"]

for row in src_kpi.iter_rows(values_only=False):
    for cell in row:
        new_cell = kpi_ws.cell(row=cell.row, column=cell.column, value=cell.value)
        if cell.has_style:
            new_cell.font = Font(name=cell.font.name, size=cell.font.size, bold=cell.font.bold, color=cell.font.color)
            new_cell.fill = PatternFill(fill_type=cell.fill.fill_type, start_color=cell.fill.start_color, end_color=cell.fill.end_color)
            new_cell.alignment = Alignment(horizontal=cell.alignment.horizontal, vertical=cell.alignment.vertical)
            new_cell.border = Border(
                left=cell.border.left, right=cell.border.right,
                top=cell.border.top, bottom=cell.border.bottom
            )
            if cell.number_format:
                new_cell.number_format = cell.number_format

kpi_ws.merge_cells("A1:G2")
for col_l, w in [('A', 36), ('B', 45), ('C', 28), ('E', 38), ('F', 14), ('G', 12), ('H', 26)]:
    kpi_ws.column_dimensions[col_l].width = w

kpi_wb.save("ML_15_Summary_KPI_Dashboard.xlsx")
print("Saved ML_15_Summary_KPI_Dashboard.xlsx")
